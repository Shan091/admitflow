"""Untrusted institutional files must extract safely or fail explicitly."""

import io
import zipfile

import pytest

from src.knowledge.extraction import ExtractionError, extract_file, validate_path


@pytest.mark.parametrize("path", ["../secret.txt", "/secret.txt", "C:\\secret.txt", "a/../x.pdf"])
def test_folder_paths_cannot_escape_the_upload_namespace(path: str) -> None:
    with pytest.raises(ExtractionError, match="unsafe_path"):
        validate_path(path)


def test_unicode_text_and_nested_folders_are_preserved() -> None:
    assert validate_path("Prospectus/fees.txt") == "Prospectus/fees.txt"
    assert extract_file("fees.txt", "ഹോസ്റ്റൽ fee 5000".encode()) == "ഹോസ്റ്റൽ fee 5000"


@pytest.mark.parametrize(
    "filename,data,code",
    [
        ("program.exe", b"not a document", "unsupported_type"),
        ("empty.txt", b"  ", "no_text"),
        ("binary.txt", b"abc\x00def", "invalid_text"),
        ("fake.pdf", b"not a pdf", "invalid_document"),
        ("fake.docx", b"not a zip", "invalid_document"),
    ],
)
def test_invalid_documents_are_not_silently_indexed(filename: str, data: bytes, code: str) -> None:
    with pytest.raises(ExtractionError, match=code):
        extract_file(filename, data)


def test_pdf_text_is_extracted_and_scans_need_ocr() -> None:
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    blank = io.BytesIO()
    writer.write(blank)
    with pytest.raises(ExtractionError, match="ocr_required"):
        extract_file("scan.pdf", blank.getvalue())
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 50 700 Td (Hostel fee is 5000 INR.) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    assert "Hostel fee is 5000 INR." in extract_file("prospectus.pdf", output.getvalue())


def test_docx_extracts_paragraphs_and_tables() -> None:
    from docx import Document

    doc = Document()
    doc.add_paragraph("Admissions are open.")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Hostel fee: 5000"
    output = io.BytesIO()
    doc.save(output)
    text = extract_file("prospectus.docx", output.getvalue())
    assert "Admissions are open." in text
    assert "Hostel fee: 5000" in text


def test_compressed_docx_bombs_are_rejected() -> None:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "A" * 2_000_000)
    with pytest.raises(ExtractionError, match="archive_limit"):
        extract_file("bomb.docx", output.getvalue())
