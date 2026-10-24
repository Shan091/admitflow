"""Section and table boundaries survive extraction and chunking."""

import io

import pytest
from docx import Document

from src.knowledge.chunking import semantic_chunks
from src.knowledge.extraction import ExtractionError, extract_file


def test_course_sections_and_table_rows_keep_their_headers() -> None:
    rows = [f"| Nursing item {i} | INR {1000 + i} |" for i in range(12)]
    text = (
        "# Synthetic Institute\n\n## Nursing\n\n"
        "| Item | Annual fee |\n| --- | --- |\n"
        + "\n".join(rows)
        + "\n\n## Commerce\n\n"
        + "Commerce students study accounting. " * 16
    )
    chunks = semantic_chunks(text, chunk_size=220, overlap=20)
    assert all(len(chunk) <= 220 for chunk in chunks)
    nursing = [chunk for chunk in chunks if "## Nursing" in chunk]
    commerce = [chunk for chunk in chunks if "## Commerce" in chunk]
    assert nursing and commerce
    for chunk in nursing:
        assert chunk.startswith("# Synthetic Institute\n## Nursing")
        assert "| Item | Annual fee |\n| --- | --- |" in chunk
        assert "Commerce" not in chunk
    for row in rows:
        assert sum(row in chunk.splitlines() for chunk in nursing) == 1
    assert all(chunk.startswith("# Synthetic Institute\n## Commerce") for chunk in commerce)
    assert all("Nursing item" not in chunk for chunk in commerce)


def test_oversized_table_structure_is_rejected_without_losing_columns() -> None:
    table = "# Course\n\n| Item | Description |\n| --- | --- |\n| A | " + "x" * 500 + " |"
    with pytest.raises(ExtractionError, match="structure_limit"):
        semantic_chunks(table, chunk_size=200, overlap=20)


@pytest.mark.parametrize("size,overlap", [(0, 0), (100, 100), (100, -1)])
def test_chunk_parameters_cannot_loop_forever(size: int, overlap: int) -> None:
    with pytest.raises(ValueError):
        semantic_chunks("Reference text", chunk_size=size, overlap=overlap)


def test_docx_preserves_heading_table_and_following_paragraph_order() -> None:
    document = Document()
    document.add_heading("Nursing", level=1)
    document.add_paragraph("The following charges are synthetic tuition examples.")
    table = document.add_table(rows=2, cols=2)
    for cell, text in zip(table.rows[0].cells, ("Item", "Tuition"), strict=True):
        cell.text = text
    for cell, text in zip(table.rows[1].cells, ("Nursing", "INR 45000"), strict=True):
        cell.text = text
    document.add_paragraph("End of the nursing section.")
    output = io.BytesIO()
    document.save(output)
    text = extract_file("prospectus.docx", output.getvalue())
    assert "# Nursing" in text
    assert "| Item | Tuition |\n| --- | --- |\n| Nursing | INR 45000 |" in text
    assert text.index("INR 45000") < text.index("End of the nursing section")
